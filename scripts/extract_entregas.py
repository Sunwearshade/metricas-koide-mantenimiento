import openpyxl
import json
import os
from datetime import datetime, date

from metricos_db import require_env, save_entregas

# Ruta en .env (usar ruta UNC \\servidor\recurso\... en el servicio de Windows)
PATH = require_env('ENTREGAS_EXCEL_PATH')

MESES = ['Ene', 'Feb', 'Mar', 'Abr', 'May', 'Jun', 'Jul', 'Ago', 'Sep', 'Oct', 'Nov', 'Dic']


def to_date(v):
    if isinstance(v, datetime):
        return v.strftime('%Y-%m-%d')
    if isinstance(v, date):
        return v.strftime('%Y-%m-%d')
    if v is None:
        return ''
    s = str(v).strip()
    if s[:1].isdigit() and len(s) >= 10:
        return s[:10]
    return s


def to_num(v):
    if v is None:
        return ''
    if isinstance(v, (int, float)):
        return int(v) if float(v).is_integer() else v
    return str(v).strip()


def main():
    wb = openpyxl.load_workbook(PATH, data_only=True)
    ws = wb['ENTREGAS']

    items = []
    for r in range(2, ws.max_row + 1):
        proveedor = str(ws.cell(r, 1).value or '').strip()
        material = str(ws.cell(r, 2).value or '').strip()
        depto = str(ws.cell(r, 4).value or '').strip()
        estatus_raw = str(ws.cell(r, 10).value or '').strip().upper()
        if not (proveedor or material):
            continue
        # Solo departamento de mantenimiento: el codigo termina en MTTO (+ numero, p.ej. ...-MTTO20)
        if not depto.upper().rstrip('0123456789. ').endswith('MTTO'):
            continue
        if estatus_raw == 'CANCELADO':
            continue

        fecha_envio = to_date(ws.cell(r, 7).value)
        fecha_estimada = to_date(ws.cell(r, 8).value)
        dias_raw = ws.cell(r, 9).value
        dias = int(dias_raw) if isinstance(dias_raw, (int, float)) else (
            str(dias_raw).strip() if dias_raw is not None else '')

        mes = -1
        if fecha_envio and len(fecha_envio) >= 7 and fecha_envio[:4].isdigit():
            try:
                mes = int(fecha_envio[5:7]) - 1
            except ValueError:
                mes = -1

        items.append({
            'proveedor': proveedor,
            'material': material,
            'cantidad': to_num(ws.cell(r, 3).value),
            'depto': depto,
            'serie': str(ws.cell(r, 5).value or '').strip(),
            'po': to_num(ws.cell(r, 6).value),
            'fecha_envio': fecha_envio,
            'fecha_estimada': fecha_estimada,
            'dias': dias,
            'estatus': estatus_raw,
            'observaciones': str(ws.cell(r, 11).value or '').strip(),
            'mes': mes,
        })

    save_entregas(items)

    entregados = sum(1 for it in items if it['estatus'] == 'ENTREGADO')
    pendientes = sum(1 for it in items if it['estatus'] in ('PENDIENTE', 'PARCIAL'))
    print('TIEMPOS DE ENTREGA (MTTO): %d partidas' % len(items))
    print('ENTREGADO: %d | PENDIENTE/PARCIAL: %d' % (entregados, pendientes))
    por_mes = {}
    for it in items:
        if it['mes'] >= 0:
            por_mes.setdefault(it['mes'], {'ent': 0, 'pend': 0, 'total': 0})
            por_mes[it['mes']]['total'] += 1
            if it['estatus'] == 'ENTREGADO':
                por_mes[it['mes']]['ent'] += 1
            else:
                por_mes[it['mes']]['pend'] += 1
    for m in sorted(por_mes):
        d = por_mes[m]
        print('  %s 2026: Total=%d | Entregado=%d | Pendiente=%d' % (MESES[m], d['total'], d['ent'], d['pend']))
    print('Guardado en MySQL (tabla entregas)')


if __name__ == '__main__':
    main()