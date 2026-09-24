"""Genera Excel sinteticos con la misma estructura que los archivos reales de Z:
para probar extract_v4.py y extract_entregas.py.

    python make_excels.py <carpeta_salida> <password>

Crea:
  MANTENIMIENTO.xlsx      (cifrado con <password>, hojas 1-ENERO y 2.FEBRERO)
  TIEMPO DE ENTREGA.xlsx  (hoja ENTREGAS)
"""
import io
import os
import sys
from datetime import datetime

import openpyxl
from msoffcrypto.format.ooxml import OOXMLFile


def mantenimiento(path, password):
    wb = openpyxl.Workbook()
    wb.remove(wb.active)
    filas = {
        '1-ENERO': [
            # cot, proveedor, f_elab, f_entrega, moneda, producto, obs, cant, unidad, precio, importe, iva%, iva, total, termino, proyecto, po, comentario
            (1, 'PROVEEDOR A', datetime(2026, 1, 5), datetime(2026, 1, 20), 'MXN', 'PLACA', 'placa guia', 2, 'PZA', 100.0, 200.0, 0.16, 32.0, 232.0, 'CONTADO', 'P-1', 15001, ''),
            (1, 'PROVEEDOR A', datetime(2026, 1, 5), '20/01/2026', 'MXN', 'TORNILLO', 'm8', 10, 'PZA', 5.5, None, 0.16, None, None, 'CONTADO', 'P-1', None, ''),
            (2, 'PROVEEDOR B', datetime(2026, 1, 9), None, 'USD', 'SENSOR', 'inductivo', 1, 'PZA', 1500, 1500, 0.16, 240, 1740, '30 DIAS', '', None, ''),
            (3, 'PROVEEDOR C', datetime(2026, 1, 12), datetime(2026, 1, 30), 'MXN', 'SERVICIO', 'rebobinado', 1, 'SERV', 8000, 8000, 0.16, 1280, 9280, '', '', 15002, 'TALLER EXTERNO'),
        ],
        '2.FEBRERO': [
            (4, 'PROVEEDOR A', datetime(2026, 2, 2), datetime(2026, 2, 15), 'MXN', 'BANDA', 'banda A-42', 3, 'PZA', 250, 750, 0.16, 120, 870, 'CONTADO', '', 15003, ''),
            ('texto', 'FILA IGNORADA', None, None, '', 'X', '', 1, '', 1, 1, 0, 0, 1, '', '', '', ''),
        ],
    }
    cols = [1, 2, 3, 4, 5, 6, 7, 8, 9, 10, 11, 12, 13, 20, 21, 24, 30, 32]
    for nombre, rows in filas.items():
        ws = wb.create_sheet(nombre)
        ws.cell(1, 1, 'REQUISICIONES MANTENIMIENTO')
        for i, row in enumerate(rows):
            for c, v in zip(cols, row):
                if v is not None:
                    ws.cell(8 + i, c, v)
    plain = io.BytesIO()
    wb.save(plain)
    plain.seek(0)
    with open(path, 'wb') as out:
        OOXMLFile(plain).encrypt(password, out)


def entregas(path):
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = 'ENTREGAS'
    ws.append(['PROVEEDOR', 'MATERIAL', 'CANT', 'DEPTO', 'SERIE', 'PO', 'ENVIO', 'ESTIMADA', 'DIAS', 'ESTATUS', 'OBS'])
    ws.append(['PROVEEDOR A', 'PLACA', 2, 'CO-COT-1-MTTO44', 'KMXZ', 15001, datetime(2026, 1, 10), datetime(2026, 1, 20), 10, 'ENTREGADO', ''])
    ws.append(['PROVEEDOR C', 'SERVICIO', 1, 'CO-COT-2-MTTO', 'KMXZ', 15002, datetime(2026, 1, 15), datetime(2026, 2, 1), 'N/A', 'PENDIENTE', 'en proceso'])
    ws.append(['PROVEEDOR A', 'BANDA', 3, 'CO-COT-3-MTTO7', 'KMXZ', 15003, datetime(2026, 2, 3), datetime(2026, 2, 15), 12, 'PARCIAL', ''])
    ws.append(['OTRO DEPTO', 'PAPEL', 5, 'CO-COT-4-PROD', 'KMXZ', 15004, datetime(2026, 2, 3), None, None, 'ENTREGADO', ''])
    ws.append(['PROVEEDOR D', 'CANCELADA', 1, 'CO-COT-5-MTTO', 'KMXZ', 15005, datetime(2026, 2, 3), None, None, 'CANCELADO', ''])
    wb.save(path)


if __name__ == '__main__':
    out_dir, password = sys.argv[1], sys.argv[2]
    os.makedirs(out_dir, exist_ok=True)
    mantenimiento(os.path.join(out_dir, 'MANTENIMIENTO.xlsx'), password)
    entregas(os.path.join(out_dir, 'TIEMPO DE ENTREGA.xlsx'))
    print('ok')
